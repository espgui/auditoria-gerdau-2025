import pandas as pd, zipfile, hashlib
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
O=Path('/mnt/data'); blue='1F4E78'
df=pd.read_csv(O/'base_sintetica_imobilizado_gerdau_2025.csv.gz',compression='gzip',parse_dates=['data_aquisicao','data_entrada_operacao'])
tests=pd.read_csv(O/'resultados_testes_base.csv'); sample=pd.read_csv(O/'amostra_pt530.csv'); truth=pd.read_csv(O/'registro_excecoes_inseridas.csv')
dictionary=[('id_bem','Texto','Identificador sintético do bem.'),('segmento','Texto','Segmento geográfico sintético.'),('pais','Texto','País compatível com o segmento.'),('unidade_sintetica','Texto','Unidade operacional artificial.'),('classe_imobilizado','Texto','Classe do ativo.'),('data_aquisicao','Data','Data simulada de aquisição.'),('data_entrada_operacao','Data','Data simulada de disponibilidade para uso.'),('custo_historico_rs','Decimal','Custo sintético em reais.'),('vida_util_anos','Inteiro','Vida útil simulada.'),('taxa_depreciacao_anual','Decimal','Taxa linear, 1/vida útil.'),('depreciacao_acumulada_rs','Decimal','Depreciação acumulada sintética.'),('impairment_acumulado_rs','Decimal','Perda acumulada sintética.'),('valor_contabil_liquido_rs','Decimal','Saldo líquido conciliado.'),('status_operacional','Texto','Ativo, ocioso, manutenção ou construção.'),('centro_custo','Texto','Centro de custo sintético.'),('projeto_capex','Texto','Projeto sintético quando obra.'),('indicador_excecao','Booleano','Marca erro proposital.'),('tipo_excecao','Texto','Natureza do erro proposital.'),('dados_sinteticos','Booleano','Sempre verdadeiro.')]
wb=Workbook();wb.remove(wb.active);nav=PatternFill('solid',fgColor=blue);white=Font(color='FFFFFF',bold=True)
def fmt(ws,widths):
 ws.sheet_view.showGridLines=False;ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.fill=nav;c.font=white;c.alignment=Alignment(horizontal='center',wrap_text=True)
 for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
 for row in ws.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
ws=wb.create_sheet('Resumo');ws.append(['Indicador','Valor','Unidade / conclusão'])
for r in [('Linhas',len(df),'registros'),('Saldo publicado',30640833000,'R$'),('Saldo base',round(df.valor_contabil_liquido_rs.sum(),2),'R$'),('Diferença','=B4-B3','R$'),('Exceções marcadas',int(df.indicador_excecao.sum()),'linhas'),('Amostra PPS',len(sample),'itens'),('Semente',20251025,'reprodução')]:ws.append(r)
fmt(ws,[35,25,38]);
for r in range(2,9):ws.cell(r,2).number_format='#,##0.00;[Red](#,##0.00);-'
ws=wb.create_sheet('Dicionario');ws.append(['Variável','Tipo','Descrição']);
for r in dictionary:ws.append(r)
fmt(ws,[32,18,70])
ws=wb.create_sheet('Excecoes');ws.append(list(tests.columns));
for r in tests.itertuples(index=False):ws.append(list(r))
fmt(ws,[45,18,18])
ws=wb.create_sheet('Amostra_200');ws.append(list(sample.columns));
for r in sample.itertuples(index=False):ws.append(list(r))
fmt(ws,[18]*len(sample.columns))
ws=wb.create_sheet('Por_Classe');ws.append(['Classe','Quantidade','Valor líquido R$'])
for k,g in df.groupby('classe_imobilizado'):ws.append([k,len(g),g.valor_contabil_liquido_rs.sum()])
fmt(ws,[32,18,25])
ws=wb.create_sheet('Por_Segmento');ws.append(['Segmento','Quantidade','Valor líquido R$'])
for k,g in df.groupby('segmento'):ws.append([k,len(g),g.valor_contabil_liquido_rs.sum()])
fmt(ws,[30,18,25])
ws=wb.create_sheet('Identificacao');
for i,(a,b) in enumerate([('Nome','______________________________________________'),('Matrícula','______________________________________________'),('Preparado por','______________________________________________'),('Data','______________________________________________'),('Revisado por','______________________________________________')],1):ws.cell(i,1,a).font=Font(bold=True);ws.cell(i,2,b)
ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=55;ws.sheet_view.showGridLines=False
wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.save(O/'Entrega_4_Base_Imobilizado_Amostragem_Gerdau_2025.xlsx')
# docx
d=Document();s=d.sections[0];s.top_margin=Inches(.7);s.bottom_margin=Inches(.7);s.left_margin=Inches(.75);s.right_margin=Inches(.75)
d.styles['Normal'].font.name='Aptos';d.styles['Normal'].font.size=Pt(10)
for n,z in [('Title',22),('Heading 1',15),('Heading 2',12)]:d.styles[n].font.name='Aptos Display';d.styles[n].font.size=Pt(z);d.styles[n].font.color.rgb=RGBColor.from_string(blue)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('ENTREGA 4\nBASE SINTÉTICA DO IMOBILIZADO E PT 530');r.bold=True;r.font.size=Pt(22);r.font.color.rgb=RGBColor.from_string(blue)
for x in ['Nome: ______________________________________________','Matrícula: __________________________________________','Preparado por: ______________________________________','Data: _______________________________________________','Revisado por: ________________________________________']:p=d.add_paragraph(x);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_heading('1. Escolha da conta',1);d.add_paragraph('O imobilizado consolidado foi escolhido por ser a maior conta individual do ativo em 31/12/2025, com saldo publicado de R$ 30.640.833 mil. A base representa 500.000 bens individuais sintéticos e não corresponde ao cadastro real da Gerdau S.A.')
d.add_heading('2. Premissas e método',1);d.add_paragraph('Geração pseudorrandômica com NumPy, semente 20251025, classes de ativos, segmentos, países, unidades, datas, vidas úteis, custo, depreciação e impairment. O valor líquido foi escalonado e ajustado para conciliar exatamente ao saldo publicado de R$ 30.640.833.000. Cada linha é um registro individual; não houve repetição de linhas para atingir o mínimo.')
d.add_heading('3. Conciliação',1);d.add_paragraph(f'Saldo publicado: R$ 30.640.833.000,00. Saldo da base: R$ {df.valor_contabil_liquido_rs.sum():,.2f}. Diferença arredondada: R$ {df.valor_contabil_liquido_rs.sum()-30640833000:,.2f}.')
d.add_heading('4. Exceções inseridas e testes',1)
t=d.add_table(rows=1,cols=4);t.style='Table Grid';
for i,x in enumerate(['Teste','Inseridas','Identificadas','Conclusão']):t.rows[0].cells[i].text=x
for r in tests.itertuples(index=False):
 c=t.add_row().cells;c[0].text=str(r.teste);c[1].text=str(r.inseridas);c[2].text=str(r.identificadas);c[3].text='Detecção integral' if r.inseridas==r.identificadas else 'Investigar'
d.add_paragraph('Foram marcadas 350 linhas de exceção. Duplicidades envolvem 120 linhas e 60 pares de identificadores. O arquivo de verdade permite comparar inserção e detecção.')
d.add_heading('5. PT 530 | Amostragem',1);d.add_paragraph('Norma: NBC TA 530. Objetivo: aplicar amostragem à população sintética do imobilizado. População: 500.000 registros; unidade amostral: bem individual; saldo: R$ 30,641 bilhões. Método: seleção com probabilidade proporcional ao valor contábil líquido, sem reposição, 200 itens. A PPS privilegia itens monetariamente relevantes. Risco amostral existe porque a amostra pode não representar todas as exceções de atributos, razão pela qual os quatro testes de integridade também foram executados sobre 100% da população.')
d.add_heading('6. Resultado e avaliação',1);d.add_paragraph('Os testes de dados completos identificaram integralmente as exceções propositalmente inseridas: 120 linhas com ID duplicado, 80 vidas úteis inconsistentes, 60 datas futuras e 90 depreciações superiores ao custo. A amostra PPS foi gerada como evidência didática, mas as exceções foram avaliadas por teste censitário, portanto não foi necessária projeção estatística das exceções conhecidas. Em auditoria real, documentos-suporte seriam examinados para cada item da amostra.')
d.add_heading('7. Limitações',1);d.add_paragraph('Os dados são integralmente sintéticos. Segmentos, unidades, custos, datas, vidas úteis e exceções não descrevem bens reais. A conciliação ao saldo publicado é matemática e não constitui evidência de existência, propriedade ou avaliação dos ativos da companhia.')
d.add_heading('8. Reprodução',1);d.add_paragraph('Executar: python gerar_base_imobilizado.py. Dependências: Python 3, NumPy e pandas. O script recria a base comprimida, o registro de exceções, os resultados dos testes e a amostra com a mesma semente.')
d.save(O/'Entrega_4_PT530_Base_Imobilizado_Gerdau_2025.docx')
readme='''# Entrega 4 - Base sintética do imobilizado\n\nNome: ______________________________________________\nMatrícula: __________________________________________\n\n## Reprodução\n`python gerar_base_imobilizado.py`\n\nDependências: Python 3, numpy e pandas. Semente: 20251025.\n\nA base possui 500.000 linhas e é integralmente sintética. O saldo líquido concilia a R$ 30.640.833.000. O arquivo CSV está comprimido em gzip e pode ser aberto com pandas ou descompactado por ferramentas padrão.\n'''
(O/'README_Entrega_4.md').write_text(readme,encoding='utf-8')
files=['base_sintetica_imobilizado_gerdau_2025.csv.gz','gerar_base_imobilizado.py','registro_excecoes_inseridas.csv','resultados_testes_base.csv','amostra_pt530.csv','Entrega_4_Base_Imobilizado_Amostragem_Gerdau_2025.xlsx','Entrega_4_PT530_Base_Imobilizado_Gerdau_2025.docx','README_Entrega_4.md']
with zipfile.ZipFile(O/'Entrega_4_Pacote_Base_Imobilizado_Gerdau_2025.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in files:z.write(O/f,arcname=f)
print(len(df),round(df.valor_contabil_liquido_rs.sum(),2),tests.to_dict('records'))
